from PySide6.QtCore import Qt
from PySide6.QtGui import QTextDocument
from PySide6.QtPrintSupport import QPrinter, QPrintDialog, QPrintPreviewDialog
from PySide6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QPushButton, QTableWidget, QTableWidgetItem,
    QHeaderView, QLabel, QMessageBox, QFileDialog
)


class TelaRelatorio(QDialog):
    """
    Janela genérica de relatório (seção 27): mostra os dados em tabela e permite
    visualizar/imprimir, exportar PDF e exportar Excel — reaproveitada por todos
    os relatórios do sistema, cada um só muda o título, as colunas e os dados.
    """

    def __init__(self, titulo: str, colunas: list[tuple[str, str]], linhas: list[dict], parent=None):
        """
        colunas: lista de (chave_no_dict, rótulo_da_coluna)
        linhas: lista de dicts com os dados (podem ter chaves a mais, só as de `colunas` são usadas)
        """
        super().__init__(parent)
        self.titulo = titulo
        self.colunas = colunas
        self.linhas = linhas

        self.setWindowTitle(titulo)
        self.resize(900, 600)
        self._montar_interface()
        self._preencher_tabela()

    def _montar_interface(self):
        layout = QVBoxLayout(self)

        cabecalho = QHBoxLayout()
        cabecalho.addWidget(QLabel(f"<b>{self.titulo}</b>"))
        cabecalho.addStretch()
        self.lbl_total = QLabel()
        cabecalho.addWidget(self.lbl_total)
        layout.addLayout(cabecalho)

        self.tabela = QTableWidget()
        self.tabela.setColumnCount(len(self.colunas))
        self.tabela.setHorizontalHeaderLabels([rotulo for _, rotulo in self.colunas])
        self.tabela.setEditTriggers(QTableWidget.NoEditTriggers)
        self.tabela.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        layout.addWidget(self.tabela)

        linha_botoes = QHBoxLayout()
        botao_imprimir = QPushButton("🖨️ Imprimir")
        botao_imprimir.clicked.connect(self._imprimir)
        botao_pdf = QPushButton("Exportar PDF")
        botao_pdf.clicked.connect(self._exportar_pdf)
        botao_excel = QPushButton("Exportar Excel")
        botao_excel.clicked.connect(self._exportar_excel)
        linha_botoes.addWidget(botao_imprimir)
        linha_botoes.addWidget(botao_pdf)
        linha_botoes.addWidget(botao_excel)
        linha_botoes.addStretch()
        layout.addLayout(linha_botoes)

    def _preencher_tabela(self):
        self.tabela.setRowCount(len(self.linhas))
        for linha_idx, registro in enumerate(self.linhas):
            for coluna_idx, (chave, _rotulo) in enumerate(self.colunas):
                valor = registro.get(chave)
                texto = "—" if valor in (None, "") else str(valor)
                self.tabela.setItem(linha_idx, coluna_idx, QTableWidgetItem(texto))
        self.lbl_total.setText(f"{len(self.linhas)} registro(s)")

    # ------------------------------------------------------------ geração HTML

    def _gerar_html(self) -> str:
        cabecalho_html = "".join(f"<th>{rotulo}</th>" for _, rotulo in self.colunas)
        linhas_html = ""
        for registro in self.linhas:
            celulas = "".join(
                f"<td>{registro.get(chave) if registro.get(chave) not in (None, '') else '—'}</td>"
                for chave, _ in self.colunas
            )
            linhas_html += f"<tr>{celulas}</tr>"

        return f"""
        <html><head><style>
            body {{ font-family: Arial, sans-serif; font-size: 10pt; }}
            h2 {{ margin-bottom: 4px; }}
            table {{ border-collapse: collapse; width: 100%; }}
            th, td {{ border: 1px solid #999; padding: 4px 6px; text-align: left; }}
            th {{ background-color: #eee; }}
        </style></head><body>
            <h2>{self.titulo}</h2>
            <p>{len(self.linhas)} registro(s)</p>
            <table><thead><tr>{cabecalho_html}</tr></thead><tbody>{linhas_html}</tbody></table>
        </body></html>
        """

    # ------------------------------------------------------------------ ações

    def _imprimir(self):
        documento = QTextDocument()
        documento.setHtml(self._gerar_html())
        impressora = QPrinter(QPrinter.HighResolution)
        dialogo = QPrintDialog(impressora, self)
        if dialogo.exec() == QDialog.Accepted:
            documento.print_(impressora)

    def _exportar_pdf(self):
        caminho, _ = QFileDialog.getSaveFileName(self, "Exportar PDF", f"{self.titulo}.pdf", "PDF (*.pdf)")
        if not caminho:
            return
        documento = QTextDocument()
        documento.setHtml(self._gerar_html())
        impressora = QPrinter(QPrinter.HighResolution)
        impressora.setOutputFormat(QPrinter.PdfFormat)
        impressora.setOutputFileName(caminho)
        documento.print_(impressora)
        QMessageBox.information(self, "Concluído", f"PDF salvo em:\n{caminho}")

    def _exportar_excel(self):
        caminho, _ = QFileDialog.getSaveFileName(self, "Exportar Excel", f"{self.titulo}.xlsx", "Excel (*.xlsx)")
        if not caminho:
            return
        import openpyxl
        pasta_trabalho = openpyxl.Workbook()
        planilha = pasta_trabalho.active
        planilha.title = "Relatório"
        planilha.append([rotulo for _, rotulo in self.colunas])
        for registro in self.linhas:
            planilha.append([registro.get(chave) or "" for chave, _ in self.colunas])
        pasta_trabalho.save(caminho)
        QMessageBox.information(self, "Concluído", f"Excel salvo em:\n{caminho}")
